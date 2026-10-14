from app.api.upload import chunk_pages
from app.llm import answer as answer_module
from app.llm.answer import NOT_FOUND, build_context, generate_answer, parse_citations
from app.loaders.document_loader import extract_pages
from app.retrieval.bm25_store import BM25Index, tokenize
from app.retrieval.hybrid import reciprocal_rank_fusion


def test_tokenize_lowercases_and_drops_stopwords_and_punctuation():
    assert tokenize("What is the BM25 score, exactly?") == ["bm25", "score", "exactly"]


def test_bm25_index_ranks_exact_term_match_first():
    index = BM25Index(["invoice number INV-2041 was paid", "the weather is sunny", "payment terms are net 30"])
    assert index.search("INV-2041", n_results=1)[0][0] == 0


def test_rrf_rewards_documents_found_by_both_retrievers():
    fused = reciprocal_rank_fusion(
        {
            "vector": [{"id": "a"}, {"id": "b"}, {"id": "c"}],
            "keyword": [{"id": "c"}, {"id": "d"}],
        },
        k=60,
    )
    assert fused[0]["id"] == "c"
    assert fused[0]["sources"] == ["vector", "keyword"]
    assert abs(fused[0]["score"] - (1 / 63 + 1 / 61)) < 1e-12


def test_parse_citations_keeps_valid_unique_numbers_in_order():
    assert parse_citations("Paris [2]. It is large [2][1]. See [9].", passage_count=3) == [2, 1]


def test_build_context_labels_passages_with_file_and_page():
    context = build_context([{"id": "x", "text": "Hello", "metadata": {"filename": "a.pdf", "page": 4}}])
    assert context.startswith("[1] (a.pdf, page 4)\nHello")


def test_generate_answer_returns_cited_sources(monkeypatch):
    monkeypatch.setattr(answer_module, "invoke_with_fallback", lambda messages, model=None: "It is 42 [2].")
    passages = [
        {"id": "p1", "text": "unrelated", "metadata": {"filename": "a.pdf", "page": 1}},
        {"id": "p2", "text": "the answer is 42", "metadata": {"filename": "b.pdf", "page": 7}},
    ]
    result = generate_answer("what is it?", passages)
    assert result["abstained"] is False
    assert [(c["filename"], c["page"]) for c in result["citations"]] == [("b.pdf", 7)]


def test_generate_answer_abstains_without_passages():
    result = generate_answer("anything?", [])
    assert result == {"answer": NOT_FOUND, "citations": [], "abstained": True}


def test_text_loader_and_chunking(tmp_path):
    path = tmp_path / "policy.txt"
    path.write_text("Leave policy.\n\nEmployees get 24 days of paid leave per year.")
    pages = extract_pages(path)
    chunks = chunk_pages(pages, "policy.txt")
    assert pages and chunks[0]["filename"] == "policy.txt" and chunks[0]["page"] == 1
    assert "24 days" in " ".join(chunk["text"] for chunk in chunks)


def test_docx_loader_reads_paragraphs_and_tables(tmp_path):
    from docx import Document

    document = Document()
    document.add_paragraph("Travel policy")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Daily limit"
    table.rows[0].cells[1].text = "INR 5000"
    path = tmp_path / "travel.docx"
    document.save(str(path))

    text = " ".join(extract_pages(path))
    assert "Travel policy" in text and "Daily limit | INR 5000" in text
