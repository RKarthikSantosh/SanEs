from pathlib import Path

from app.loaders.pdf_loader import extract_pdf_pages

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# Text files are split into blocks of this many chars, each treated as a page
_BLOCK_CHARS = 3000


def _group_into_blocks(paragraphs: list[str]) -> list[str]:
    blocks, current = [], ""
    for paragraph in paragraphs:
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if current and len(current) + len(paragraph) > _BLOCK_CHARS:
            blocks.append(current)
            current = ""
        current = f"{current}\n{paragraph}" if current else paragraph
    if current:
        blocks.append(current)
    return blocks


def extract_pages(file_path: Path) -> list[str]:
    """Return the text of a document as a list of pages (1-based page = index + 1)."""
    extension = file_path.suffix.lower()

    if extension == ".pdf":
        pages, _ = extract_pdf_pages(file_path)
        return pages

    if extension == ".docx":
        from docx import Document

        document = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.append(" | ".join(cell.text for cell in row.cells))
        return _group_into_blocks(paragraphs)

    if extension in {".txt", ".md"}:
        text = file_path.read_text(encoding="utf-8", errors="ignore")
        return _group_into_blocks(text.split("\n\n"))

    raise ValueError(f"Unsupported file type: {extension}")
